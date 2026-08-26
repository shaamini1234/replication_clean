#!/usr/bin/env python3
"""build_handover_doc.py — render newhandover.md as newhandover.docx.

The markdown file is the single source of truth; this produces the Word version
for circulation. Run after editing the handover or redrawing its charts:

    ./venv/bin/python python/build_handover_doc.py

Handles the constructs the handover uses: headings, paragraphs with bold and
inline code, bullet lists, tables, fenced code blocks, and images with italic
captions.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent
SRC = REPO / "newhandover.md"
DEST = REPO / "newhandover.docx"

INK = RGBColor(0x1A, 0x1C, 0x22)
MUT = RGBColor(0x5C, 0x62, 0x70)
NAVY = RGBColor(0x1F, 0x38, 0x64)


def setup(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, colour in (("Title", 22, NAVY), ("Heading 1", 15, NAVY),
                               ("Heading 2", 12.5, NAVY), ("Heading 3", 11, INK)):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.color.rgb = colour
        st.font.bold = name != "Title"
        st.paragraph_format.space_before = Pt(14 if name.startswith("Heading") else 0)
        st.paragraph_format.space_after = Pt(6)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def write_runs(par, text: str) -> None:
    """Render **bold**, `code` and *italic* inside a paragraph."""
    text = text.replace("&nbsp;", " ")
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = par.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            write_runs(par, text)
            for run in par.runs:
                run.font.size = Pt(9.5)
                if r == 0:
                    run.bold = True
    doc.add_paragraph()


def main() -> int:
    if not SRC.exists():
        sys.exit(f"Missing {SRC.name}")
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = section.bottom_margin = Inches(0.9)
    setup(doc)

    lines = SRC.read_text().splitlines()
    i, table_buf, code_buf, in_code = 0, [], [], False
    images = 0

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, table_buf)
            table_buf = []

    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("```"):
            if in_code:
                par = doc.add_paragraph()
                run = par.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                par.paragraph_format.left_indent = Inches(0.25)
                code_buf, in_code = [], False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        img = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if img:
            flush_table()
            path = REPO / img.group(2)
            if path.exists():
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(path), width=Inches(6.2))
                images += 1
            i += 1
            continue

        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if not all(set(c) <= set("-: ") for c in cells):
                table_buf.append(cells)
            i += 1
            continue
        flush_table()

        if not line:
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            text = line.lstrip("# ").strip().replace("`", "")  # headings take no code marks
            if level == 1:
                doc.add_paragraph(text, style="Title")
            else:
                doc.add_paragraph(text, style=f"Heading {min(level - 1, 3)}")
            i += 1
            continue
        if re.match(r"^[-*] ", line):
            par = doc.add_paragraph(style="List Bullet")
            write_runs(par, line[2:])
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            par = doc.add_paragraph(style="List Number")
            write_runs(par, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue

        par = doc.add_paragraph()
        write_runs(par, line)
        # an italic-only line is a figure caption
        if line.startswith("*") and line.endswith("*"):
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in par.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = MUT
        i += 1

    flush_table()
    doc.save(DEST)
    print(f"wrote {DEST.name}: {len(doc.paragraphs)} paragraphs, "
          f"{len(doc.tables)} tables, {images} images")
    return 0


if __name__ == "__main__":
    sys.exit(main())
